from __future__ import annotations

import argparse
import csv
import re
import xml.etree.ElementTree as ET
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable
import datetime

import openpyxl
from openpyxl.chart.line_chart import LineChart
from openpyxl.utils.datetime import MAC_EPOCH, WINDOWS_EPOCH, from_excel
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


_RE_FLOW_ZONE_STRAND = re.compile(
    r"^(?P<flow>[1-7])ST_二冷水_(?P<zone>[1-7])区(?:_(?P<strand>侧弧|内弧|外弧|内外弧))?_",
)
_RE_FLOW_ZONE_STRAND_NO_TRAIL = re.compile(
    r"^(?P<flow>[1-7])ST_二冷水_(?P<zone>[1-7])区(?:_(?P<strand>侧弧|内弧|外弧|内外弧))?$",
)


def infer_circuit_key_from_stem(stem: str) -> str | None:
    s = stem.replace(" ", "")
    # Prefer the "no trailing underscore" pattern first, because many filenames
    # end exactly at the strand token (no extra '_' suffix).
    m = _RE_FLOW_ZONE_STRAND_NO_TRAIL.match(s)
    if not m:
        m = _RE_FLOW_ZONE_STRAND.match(s)
    if not m:
        return None

    flow = m.group("flow")
    zone = m.group("zone")
    strand = m.groupdict().get("strand")
    if strand:
        return f"{flow}ST_二冷水_{zone}区_{strand}"
    return f"{flow}ST_二冷水_{zone}区"


def _num_cache_values(data_source) -> list[float] | None:
    """
    openpyxl chart data sources store plot points in numCache/strCache.
    We extract cached point values so we don't depend on referenced cells staying accessible.
    """
    if data_source is None:
        return None
    cache = getattr(data_source, "numCache", None) or getattr(data_source, "strCache", None)
    if cache is None or not getattr(cache, "pt", None):
        return None
    out: list[float] = []
    for pt in cache.pt:
        v = getattr(pt, "v", None)
        if v is None:
            out.append(float("nan"))
        else:
            out.append(float(v))
    return out


def _read_numref_range_values(wb: openpyxl.Workbook, sheet_name: str, range_str: str) -> list[float]:
    """
    Read all values from a chart-referenced numeric range (e.g. $B$2:$B$3851).
    Returns a list aligned by row from start->end.
    """
    ws = wb[sheet_name]
    # openpyxl accepts e.g. "$B$2:$B$3851" directly
    cells = ws[range_str]
    # cells shape: rows x cols. We assume chart uses single column ranges.
    out: list[object] = []
    for r in cells:
        c = r[0]
        v = c.value
        if v is None:
            out.append(float("nan"))
            continue
        # start time may already be decoded as datetime
        if isinstance(v, (datetime.datetime, datetime.date)):
            out.append(v)
            continue
        if isinstance(v, str):
            if not v.strip():
                out.append(float("nan"))
                continue
            try:
                out.append(float(v))
                continue
            except Exception:
                out.append(float("nan"))
                continue
        out.append(float(v))
    # caller will decide how to interpret each element
    return out  # type: ignore[return-value]


_CHART_NS = "http://schemas.openxmlformats.org/drawingml/2006/chart"


def _chart_q(local: str) -> str:
    return f"{{{_CHART_NS}}}{local}"


def _numcache_pt_values(cache_el: ET.Element) -> list[float] | None:
    pts = cache_el.findall(_chart_q("pt"))
    if not pts:
        return None
    by_idx: dict[int, float] = {}
    for pt in pts:
        idx = int(pt.get("idx", "0"))
        vel = pt.find(_chart_q("v"))
        if vel is None or vel.text is None or not str(vel.text).strip():
            continue
        try:
            by_idx[idx] = float(vel.text)
        except ValueError:
            continue
    if not by_idx:
        return None
    m = max(by_idx)
    return [by_idx.get(i, float("nan")) for i in range(m + 1)]


def _ser_axis_values(ser: ET.Element, axis: str, wb: openpyxl.Workbook) -> list[object]:
    """
    Read category (cat) or value (val) axis for one c:ser.
    Prefers embedded numCache (折线图实际采样点) over full sheet column.
    """
    q = _chart_q
    el = ser.find(q(axis))
    if el is None:
        return []
    num_ref = el.find(q("numRef"))
    if num_ref is None:
        return []
    cache = num_ref.find(q("numCache"))
    if cache is not None:
        cached = _numcache_pt_values(cache)
        if cached is not None:
            return cached  # type: ignore[return-value]
    f_el = num_ref.find(q("f"))
    if f_el is None or not f_el.text:
        return []
    parsed = _parse_numref_f(wb, f_el.text.strip())
    if not parsed:
        return []
    sheet, rng = parsed
    return _read_numref_range_values(wb, sheet, rng)


def _extract_line_series_from_chart_xml(xlsx_path: Path, wb: openpyxl.Workbook) -> list[tuple[list[object], list[float]]] | None:
    """
    Excel 将「流量/压力」与「阀位」分在两个 c:lineChart 中时，openpyxl 顶层 LineChart.series 只有 3 条。
    这里按 OOXML 顺序收集 plotArea 下所有 lineChart 的 c:ser，从而拿到第 4 条阀位折线。
    """
    series_pack: list[tuple[list[object], list[float]]] = []
    try:
        with zipfile.ZipFile(xlsx_path, "r") as zf:
            chart_names = sorted(n for n in zf.namelist() if n.startswith("xl/charts/chart") and n.endswith(".xml"))
            for cn in chart_names:
                try:
                    root = ET.fromstring(zf.read(cn))
                except ET.ParseError:
                    continue
                for plot in root.iter(_chart_q("plotArea")):
                    for lc in plot.findall(_chart_q("lineChart")):
                        for ser in lc.findall(_chart_q("ser")):
                            xs = _ser_axis_values(ser, "cat", wb)
                            ys_raw = _ser_axis_values(ser, "val", wb)
                            ys: list[float] = []
                            for v in ys_raw:
                                if isinstance(v, (float, int)):
                                    ys.append(float(v))
                                else:
                                    try:
                                        ys.append(float(v))
                                    except Exception:
                                        ys.append(float("nan"))
                            if not xs or not ys:
                                continue
                            n = min(len(xs), len(ys))
                            if n <= 0:
                                continue
                            series_pack.append((xs[:n], ys[:n]))
    except (OSError, zipfile.BadZipFile, KeyError):
        return None
    if len(series_pack) < 3:
        return None
    n_align = min(min(len(a[0]), len(a[1])) for a in series_pack)
    if n_align < 3:
        return None
    return [(xs[:n_align], ys[:n_align]) for xs, ys in series_pack]


def _parse_numref_f(wb: openpyxl.Workbook, numRef_f: str) -> tuple[str, str] | None:
    """
    numRef.f format in exported charts:
      'Sheet1!$B$2:$B$3851'
    """
    if "!" not in numRef_f:
        return None
    sheet_name, rng = numRef_f.split("!", 1)
    sheet_name = sheet_name.strip("'\"")
    rng = rng.strip()
    if sheet_name not in wb.sheetnames:
        # sometimes openpyxl leaves sheet name as a quoted string
        return None
    return sheet_name, rng


@dataclass(frozen=True)
class ExportSpec:
    xlsx_dir: Path
    out_data_dir: Path
    out_tag_catalog_csv: Path
    synthetic_iot_base: int = 1000000000000
    max_files_debug: int | None = None


def export_xlsx_charts_to_synthetic_ot_csv(spec: ExportSpec) -> Path:
    spec.out_data_dir.mkdir(parents=True, exist_ok=True)
    rows_catalog: list[dict[str, str]] = []

    xlsx_files = sorted(p for p in spec.xlsx_dir.glob("*.xlsx") if not p.name.startswith("~$"))
    if spec.max_files_debug is not None:
        xlsx_files = xlsx_files[: int(spec.max_files_debug)]

    # 优先从 chart XML 收集全部 c:lineChart 下的序列（阀位常在第二个 lineChart，openpyxl 读不到）；
    # 否则回退到 openpyxl LineChart.series。
    # 约定：序列 0=流量目标，1=流量实际，2=压力，3=阀位（可选）。
    next_iot = int(spec.synthetic_iot_base)
    circuit_to_tag: dict[str, dict[str, str]] = {}

    for idx, xl_path in enumerate(xlsx_files):
        stem = xl_path.stem
        circuit_key = infer_circuit_key_from_stem(stem)
        if not circuit_key:
            # can't parse circuit membership from file name
            continue

        wb = openpyxl.load_workbook(xl_path, data_only=True)
        date1904 = False
        try:
            # openpyxl uses an absolute base date; 1904 base date means Excel 1904 system.
            date1904 = bool(wb.excel_base_date.year == 1904)
        except Exception:
            date1904 = False

        xml_pack = _extract_line_series_from_chart_xml(xl_path, wb)
        x_vals: list[object]
        y0_vals: list[float]
        y1_vals: list[float]
        y2_vals: list[float]
        y3_vals: list[float] | None

        if xml_pack and len(xml_pack) >= 3:
            x_vals = list(xml_pack[0][0])
            y0_vals = list(xml_pack[0][1])
            y1_vals = list(xml_pack[1][1])
            y2_vals = list(xml_pack[2][1])
            y3_vals = list(xml_pack[3][1]) if len(xml_pack) >= 4 else None
        else:
            chart = None
            for ws in wb.worksheets:
                for ch in getattr(ws, "_charts", []) or []:
                    if isinstance(ch, LineChart):
                        chart = ch
                        break
                if chart is not None:
                    break
            if chart is None or len(chart.series) < 3:
                wb.close()
                continue

            s0, s1, s2 = chart.series[0], chart.series[1], chart.series[2]
            s3 = chart.series[3] if len(chart.series) >= 4 else None

            def series_xy(series, wb_inner: openpyxl.Workbook, idx: int) -> tuple[list[object], list[float]]:
                cat_ds = getattr(series, "cat", None)
                val_ds = getattr(series, "val", None)

                x_cache = _num_cache_values(cat_ds)
                y_cache = _num_cache_values(val_ds)

                if x_cache is not None and y_cache is not None and len(x_cache) and len(y_cache):
                    n = min(len(x_cache), len(y_cache))
                    return x_cache[:n], y_cache[:n]  # type: ignore[return-value]

                cat_numRef = getattr(cat_ds, "numRef", None)
                val_numRef = getattr(val_ds, "numRef", None)
                if cat_numRef is None or val_numRef is None:
                    return [], []

                cat_parsed = _parse_numref_f(wb_inner, getattr(cat_numRef, "f", ""))
                val_parsed = _parse_numref_f(wb_inner, getattr(val_numRef, "f", ""))
                if not cat_parsed or not val_parsed:
                    return [], []
                sheet_x, range_x = cat_parsed
                sheet_y, range_y = val_parsed
                x_raw = _read_numref_range_values(wb_inner, sheet_x, range_x)
                y_raw = _read_numref_range_values(wb_inner, sheet_y, range_y)
                n = min(len(x_raw), len(y_raw))
                x_raw = x_raw[:n]
                y_raw = y_raw[:n]
                yr: list[float] = []
                for v in y_raw:
                    if isinstance(v, (float, int)):
                        yr.append(float(v))
                    else:
                        try:
                            yr.append(float(v))
                        except Exception:
                            yr.append(float("nan"))
                return x_raw, yr

            x_vals, y0_vals = series_xy(s0, wb, 0)
            _, y1_vals = series_xy(s1, wb, 1)
            _, y2_vals = series_xy(s2, wb, 2)
            y3_vals = None
            if s3 is not None:
                _, y3_try = series_xy(s3, wb, 3)
                y3_vals = y3_try if y3_try else None

        wb.close()

        if not x_vals or not y0_vals or not y1_vals or not y2_vals:
            continue

        n = min(len(x_vals), len(y0_vals), len(y1_vals), len(y2_vals))
        if y3_vals:
            n = min(n, len(y3_vals))
        x_vals = x_vals[:n]
        y0_vals = y0_vals[:n]
        y1_vals = y1_vals[:n]
        y2_vals = y2_vals[:n]
        if y3_vals:
            y3_vals = y3_vals[:n]

        # filter invalid points: require valid TRG/ACT/压力；若含阀位则该点也需有效
        def is_valid_x(v: object) -> bool:
            if isinstance(v, (datetime.datetime, datetime.date)):
                return True
            if isinstance(v, (float, int)):
                fv = float(v)
                return fv == fv and fv not in (float("inf"), float("-inf"))
            return False

        def is_valid_y(v: object) -> bool:
            if isinstance(v, (float, int)):
                fv = float(v)
                return fv == fv and fv not in (float("inf"), float("-inf"))
            return False

        def row_ok(i: int) -> bool:
            base = (
                is_valid_x(x_vals[i])
                and is_valid_y(y0_vals[i])
                and is_valid_y(y1_vals[i])
                and is_valid_y(y2_vals[i])
            )
            if not base:
                return False
            if y3_vals is not None:
                return is_valid_y(y3_vals[i])
            return True

        keep_idx = [i for i in range(n) if row_ok(i)]
        if len(keep_idx) < 3:
            continue

        x_vals = [x_vals[i] for i in keep_idx]
        y0_vals = [y0_vals[i] for i in keep_idx]
        y1_vals = [y1_vals[i] for i in keep_idx]
        y2_vals = [y2_vals[i] for i in keep_idx]
        if y3_vals is not None:
            y3_vals = [y3_vals[i] for i in keep_idx]

        def to_ts_str(x: object) -> str:
            if isinstance(x, datetime.datetime):
                return x.strftime("%Y-%m-%d %H:%M:%S")
            if isinstance(x, datetime.date):
                return datetime.datetime(x.year, x.month, x.day).strftime("%Y-%m-%d %H:%M:%S")
            # Excel serial date -> datetime（openpyxl 3.x 使用 epoch= 而非 date1904=）
            epoch = MAC_EPOCH if date1904 else WINDOWS_EPOCH
            dt = from_excel(float(x), epoch=epoch)
            return dt.strftime("%Y-%m-%d %H:%M:%S")

        ts_list = [to_ts_str(x) for x in x_vals]

        trg_iot = str(next_iot)
        act_iot = str(next_iot + 1)
        pres_iot = str(next_iot + 2)
        next_iot += 3
        valve_iot: str | None = None
        if y3_vals is not None:
            valve_iot = str(next_iot)
            next_iot += 1

        # write csv helper
        def write_csv(iot: str, values: list[float]) -> None:
            out_csv = spec.out_data_dir / f"{iot}.csv"
            with out_csv.open("w", encoding="utf-8-sig", newline="") as f:
                w = csv.writer(f)
                w.writerow(["ts", "value"])
                for ts, v in zip(ts_list, values):
                    w.writerow([ts, v])

        write_csv(trg_iot, y0_vals)  # series[0] => TRG
        write_csv(act_iot, y1_vals)  # series[1] => ACT
        write_csv(pres_iot, y2_vals)  # series[2] => pressure
        if valve_iot is not None and y3_vals is not None:
            write_csv(valve_iot, y3_vals)  # series[3] => 阀位

        # tag catalog rows: stage1 infers kind from name substring
        tag_catalog: list[tuple[str, str]] = [
            (trg_iot, f"{circuit_key}_流量_TRG"),
            (act_iot, f"{circuit_key}_流量_ACT"),
            (pres_iot, f"{circuit_key}_压力"),
        ]
        if valve_iot is not None:
            tag_catalog.append((valve_iot, f"{circuit_key}_阀位"))
        for iot, name in tag_catalog:
            rows_catalog.append({"iot": iot, "name": name})

        ctags: dict[str, str] = {
            "flow_trg_tag": trg_iot,
            "flow_act_tag": act_iot,
            "pressure_tag": pres_iot,
        }
        if valve_iot is not None:
            ctags["valve_pos_tag"] = valve_iot
        circuit_to_tag[circuit_key] = ctags

    # save catalog
    spec.out_tag_catalog_csv.parent.mkdir(parents=True, exist_ok=True)
    # minimal columns: iot,name (isa is optional)
    with spec.out_tag_catalog_csv.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["iot", "name"])
        w.writeheader()
        for r in rows_catalog:
            w.writerow(r)

    # optional debug: circuit->tag mapping
    (spec.out_data_dir / "xlsx_export_circuit_to_tag.json").write_text(
        __import__("json").dumps(circuit_to_tag, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )

    return spec.out_tag_catalog_csv


def build_argparser() -> argparse.ArgumentParser:
    p = argparse.ArgumentParser(prog="xlsx_chart_export")
    p.add_argument("--xlsx_dir", required=True, help="Directory containing chart xlsx files.")
    p.add_argument("--out_data_dir", required=True, help="Output synthetic OT CSV dir.")
    p.add_argument("--out_tag_catalog_csv", required=True, help="Output synthetic tag_catalog.csv (iot->name).")
    p.add_argument("--synthetic_iot_base", type=int, default=1000000000000)
    p.add_argument("--max_files_debug", type=int, default=None)
    return p


def main() -> None:
    p = build_argparser()
    args = p.parse_args()
    spec = ExportSpec(
        xlsx_dir=Path(args.xlsx_dir),
        out_data_dir=Path(args.out_data_dir),
        out_tag_catalog_csv=Path(args.out_tag_catalog_csv),
        synthetic_iot_base=int(args.synthetic_iot_base),
        max_files_debug=int(args.max_files_debug) if args.max_files_debug else None,
    )
    out = export_xlsx_charts_to_synthetic_ot_csv(spec)
    print("OK. Wrote tag_catalog to:", out)


if __name__ == "__main__":
    main()

